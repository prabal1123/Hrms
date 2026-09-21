"""
Bulk import engine: read a sheet, validate every row, report what would happen.

Nothing here writes to the database. `parse_upload()` / `analyze_rows()` return
an ImportReport that the preview screen renders and the confirm step (added
separately) commits.

Row statuses
  ERROR    blocked: a required value is missing/invalid, or an Aadhar/PAN clashes
  SKIP     duplicate person: email already in this file, an existing employee in
           this organization, or a login that is already part of this
           organization (or that belongs to no organization at all)
  WARNING  importable, but an optional value was invalid and has been dropped
  READY    clean

A login that exists in ANOTHER organization is not a duplicate: the person is
imported here too and linked to that same login (the row shows an info note).

Only rows that will actually be imported (READY / WARNING) register their
email / Aadhar / PAN / phone, so an erroneous row never "uses up" an identifier.
"""

import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.db.models.functions import Lower
from openpyxl import load_workbook

from employees.models import Employee
from organizations.models import OrganizationMember

from .cleaning import PHONE_MIN_DIGITS, clean_value
from .columns import build_columns, find_header_row, match_headers, HEADER_SCAN_ROWS

logger = logging.getLogger(__name__)
User = get_user_model()

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 500
MAX_SCANNED_ROWS = 20000  # guards against sheets with huge blocks of formatted empty rows

READY, WARNING, SKIP, ERROR = "READY", "WARNING", "SKIP", "ERROR"


@dataclass
class RowResult:
    row_number: int  # the row number as shown in Excel
    values: dict = field(default_factory=dict)  # column.key -> what the preview shows
    data: dict = field(default_factory=dict)  # Employee field -> cleaned value (for the commit)
    checklist: dict = field(default_factory=dict)  # checklist item id -> completion date
    errors: list = field(default_factory=list)  # block the row
    skips: list = field(default_factory=list)  # duplicate person: row is skipped
    warnings: list = field(default_factory=list)  # informational; row still imports
    infos: list = field(default_factory=list)  # shown in Notes; never change the status

    @property
    def status(self):
        if self.errors:
            return ERROR
        if self.skips:
            return SKIP
        if self.warnings:
            return WARNING
        return READY

    @property
    def importable(self):
        return self.status in (READY, WARNING)

    @property
    def notes(self):
        return self.errors + self.skips + self.warnings + self.infos


@dataclass
class ImportReport:
    columns: list = field(default_factory=list)  # columns present in the sheet, in display order
    rows: list = field(default_factory=list)
    file_errors: list = field(default_factory=list)  # fatal: nothing can be imported
    notices: list = field(default_factory=list)  # non-fatal remarks about the sheet itself

    @property
    def ok(self):
        return not self.file_errors

    @property
    def importable_rows(self):
        return [row for row in self.rows if row.importable]

    @property
    def counts(self):
        counts = {READY: 0, WARNING: 0, SKIP: 0, ERROR: 0}
        for row in self.rows:
            counts[row.status] += 1
        return counts


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def read_workbook_rows(file_obj):
    """Return (rows, error). `rows` is a list of lists of raw cell values."""
    size = getattr(file_obj, "size", None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        return None, f"The file is too large (maximum {MAX_UPLOAD_BYTES // (1024 * 1024)} MB)."
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:  # openpyxl raises many types (BadZipFile, InvalidFileException, ...)
        logger.exception("Bulk import: could not open workbook")
        return None, "Could not read this file. Please upload a valid .xlsx Excel file."
    try:
        sheet = workbook.active
        if sheet is None:
            return None, "The workbook has no sheets."
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_SCANNED_ROWS:
                return None, "The sheet is too large to import."
            rows.append(list(row))
        return rows, None
    except Exception:
        logger.exception("Bulk import: could not read sheet")
        return None, "Could not read this file. Please upload a valid .xlsx Excel file."
    finally:
        workbook.close()


def parse_upload(file_obj, organization):
    """Read an uploaded .xlsx and analyze it. Always returns an ImportReport."""
    rows, error = read_workbook_rows(file_obj)
    if error:
        return ImportReport(file_errors=[error])
    return analyze_rows(rows, organization)


# ---------------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------------

def _is_blank(row):
    return not any(cell is not None and str(cell).strip() for cell in row)


def _display(raw):
    if raw is None:
        return ""
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


def _phone_key(phone):
    digits = re.sub(r"\D", "", phone or "")
    return digits[-PHONE_MIN_DIGITS:] if len(digits) >= PHONE_MIN_DIGITS else ""


def analyze_rows(rows, organization):
    report = ImportReport()
    columns = build_columns(organization)

    header_index = find_header_row(rows, columns)
    if header_index is None:
        required = ", ".join(c.label for c in columns if c.required)
        report.file_errors.append(
            f"Could not find the header row in the first {HEADER_SCAN_ROWS} rows. "
            f"Required columns: {required}."
        )
        return report

    positions, ignored, duplicates = match_headers(rows[header_index], columns)
    report.columns = [c for c in columns if c.key in positions]
    if ignored:
        report.notices.append("Ignored columns (not recognised): " + ", ".join(ignored) + ".")
    for label in duplicates:
        report.notices.append(f"Column '{label}' appears more than once; only the first is used.")

    data_rows = [
        (number, row)
        for number, row in enumerate(rows, start=1)
        if number > header_index + 1 and not _is_blank(row)
    ]
    if not data_rows:
        report.file_errors.append("No employee rows were found below the header.")
        return report
    if len(data_rows) > MAX_DATA_ROWS:
        report.file_errors.append(
            f"The sheet has {len(data_rows)} employee rows; the maximum per import is "
            f"{MAX_DATA_ROWS}. Please split the file."
        )
        return report

    results = [_clean_row(number, row, report.columns, positions) for number, row in data_rows]
    _apply_duplicate_checks(results, organization)
    report.rows = results
    return report


def _clean_row(number, row, columns, positions):
    result = RowResult(row_number=number)
    for column in columns:
        index = positions[column.key]
        raw = row[index] if index < len(row) else None
        value, error = clean_value(column, raw)

        if error is None and value is None and column.required:
            error = f"{column.display_name} is required."

        # Preview shows the cleaned value; for empty/invalid cells (and the
        # split full name) it shows what was in the sheet.
        if error or value is None or isinstance(value, dict):
            result.values[column.key] = _display(raw)
        else:
            result.values[column.key] = value

        if error:
            if column.required:
                result.errors.append(error)
            else:
                result.warnings.append(f"{error} Skipped; the row is still imported.")
            continue
        if value is None:
            continue

        if column.kind == "full_name":
            result.data.update(value)
        elif column.is_checklist:
            result.checklist[column.item_id] = value
        else:
            result.data[column.key] = value
    return result


# ---------------------------------------------------------------------------
# Duplicate detection (fixed number of queries, independent of row count)
# ---------------------------------------------------------------------------

def _existing_logins(emails):
    """
    email -> id of the User whose email or username matches it (any organization).
    The value is None when the email matches more than one account.
    """
    if not emails:
        return {}
    wanted = list(emails)
    rows = list(
        User.objects.annotate(e=Lower("email")).filter(e__in=wanted).values_list("e", "id")
    )
    rows += list(
        User.objects.annotate(u=Lower("username")).filter(u__in=wanted).values_list("u", "id")
    )
    logins = {}
    for key, user_id in rows:
        if logins.get(key, user_id) != user_id:
            logins[key] = None
        else:
            logins[key] = user_id
    return logins


def _org_memberships(user_ids, organization):
    """(ids of users in ANY organization, ids of users in THIS organization)."""
    in_any, in_this = set(), set()
    if not user_ids:
        return in_any, in_this
    rows = OrganizationMember.objects.filter(user_id__in=list(user_ids)).values_list(
        "user_id", "organization_id"
    )
    for user_id, organization_id in rows:
        in_any.add(user_id)
        if organization_id == organization.id:
            in_this.add(user_id)
    return in_any, in_this


class OrganizationIndex:
    """Identifiers of the organization's existing employees, compared in Python (they're encrypted)."""

    def __init__(self, organization):
        self.email, self.aadhar_no, self.pan_no, self.phone = {}, {}, {}, {}
        self.user_ids = set()  # logins that already have an employee record here
        employees = Employee.objects.filter(organization=organization).only(
            "employee_id", "email", "phone", "aadhar_no", "pan_no", "user_id"
        )
        for employee in employees:
            if employee.user_id:
                self.user_ids.add(employee.user_id)
            email = (employee.email or "").strip().lower()
            if email:
                self.email[email] = employee.employee_id
            if employee.aadhar_no:
                self.aadhar_no[employee.aadhar_no] = employee.employee_id
            if employee.pan_no:
                self.pan_no[employee.pan_no.upper()] = employee.employee_id
            phone = _phone_key(employee.phone)
            if phone:
                self.phone.setdefault(phone, employee.employee_id)


def _apply_duplicate_checks(results, organization):
    candidates = [r for r in results if not r.errors and r.data.get("email")]
    logins = _existing_logins({r.data["email"] for r in candidates})
    in_any_org, in_this_org = _org_memberships(
        {user_id for user_id in logins.values() if user_id is not None}, organization
    )
    existing = OrganizationIndex(organization)

    seen = {"email": {}, "aadhar_no": {}, "pan_no": {}, "phone": {}}

    for result in results:
        if result.errors:
            continue
        email = result.data["email"]

        if email in seen["email"]:
            result.skips.append(
                f"Duplicate of row {seen['email'][email]} in this file. Will be skipped."
            )
            continue
        if email in existing.email:
            result.skips.append(
                f"An employee with this email already exists ({existing.email[email]}). "
                "Will be skipped."
            )
            continue

        if email in logins:
            user_id = logins[email]
            if user_id is None:
                result.skips.append(
                    "This email matches more than one account. Will be skipped."
                )
                continue
            if user_id in in_this_org or user_id in existing.user_ids:
                result.skips.append(
                    "This person is already in this organization. Will be skipped."
                )
                continue
            if user_id not in in_any_org:
                result.skips.append(
                    "An account with this email exists but isn't part of any organization, "
                    "so it can't be added. Will be skipped."
                )
                continue
            result.infos.append(
                "Existing login (from another organization): it will be linked to this "
                "organization. Their password stays as it is."
            )

        for key, label in (("aadhar_no", "Aadhar No."), ("pan_no", "PAN No.")):
            value = result.data.get(key)
            if not value:
                continue
            if value in seen[key]:
                result.errors.append(f"{label} is also on row {seen[key][value]} of this file.")
            elif value in getattr(existing, key):
                result.errors.append(
                    f"{label} already belongs to employee {getattr(existing, key)[value]}."
                )
        if result.errors:
            continue

        phone = _phone_key(result.data.get("phone"))
        if phone:
            if phone in seen["phone"]:
                result.warnings.append(
                    f"Mobile No. is also on row {seen['phone'][phone]}; "
                    "logging in by phone won't work for either person."
                )
            elif phone in existing.phone:
                result.warnings.append(
                    f"Mobile No. is already used by employee {existing.phone[phone]}; "
                    "logging in by phone won't work for either person."
                )

        # Only importable rows register their identifiers.
        seen["email"][email] = result.row_number
        for key in ("aadhar_no", "pan_no"):
            if result.data.get(key):
                seen[key][result.data[key]] = result.row_number
        if phone:
            seen["phone"].setdefault(phone, result.row_number)