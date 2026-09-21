"""Downloadable .xlsx files: the blank template and the "rows with problems" report."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .columns import DATE
from .engine import ERROR, SKIP


def _put(sheet, row, column, value):
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # never let a value like "=1+1" become a formula
    return cell


def _to_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_template(columns):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    for number, column in enumerate(columns, start=1):
        _put(sheet, 1, number, column.label).font = Font(bold=True)
        sheet.column_dimensions[sheet.cell(row=1, column=number).column_letter].width = max(
            16, len(column.label) + 4
        )
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("Instructions")
    _put(guide, 1, 1, "Column").font = Font(bold=True)
    _put(guide, 1, 2, "Required").font = Font(bold=True)
    _put(guide, 1, 3, "Notes").font = Font(bold=True)
    for number, column in enumerate(columns, start=2):
        _put(guide, number, 1, column.label)
        _put(guide, number, 2, "Yes" if column.required else "")
        note = "Date, e.g. 31/03/2026 (a filled date marks the checklist item Done)" if (
            column.is_checklist
        ) else "Date, e.g. 31/03/2026" if column.kind == DATE else ""
        _put(guide, number, 3, note)
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["C"].width = 60
    return _to_bytes(workbook)


def build_problem_report(payload):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rows with problems"
    headers = ["Row", "Status"] + [c["name"] for c in payload["columns"]] + ["Notes"]
    for number, header in enumerate(headers, start=1):
        _put(sheet, 1, number, header).font = Font(bold=True)
    line = 2
    for row in payload["rows"]:
        if row["status"] not in (ERROR, SKIP):
            continue
        cells = [row["n"], row["status"]]
        cells += [row["values"].get(c["key"], "") for c in payload["columns"]]
        cells.append(" | ".join(row["notes"]))
        for number, value in enumerate(cells, start=1):
            _put(sheet, line, number, value)
        line += 1
    return _to_bytes(workbook)
