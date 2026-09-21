from datetime import date, datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.services import GrantLoginError, link_existing_login
from employees.bulk_import import ERROR, READY, SKIP, WARNING, analyze_rows, parse_upload
from employees.models import ChecklistItem, Employee, EmployeeChecklistEntry
from organizations.models import Organization, OrganizationMember

User = get_user_model()

HEADER = ["Emp Name", "Mail ID", "Mobile No.", "Designation", "Aadhar No.", "PAN NO.",
          "NAME OF BANK", "Bank A/C No.", "IFSC No.", "Emp Adress", "Date of Joining",
          "Date of AVSEC Training", "Date of Police Verification", "UAN No.", "ESIC No."]
KEYS = ["name", "mail", "phone", "desig", "aadhar", "pan", "bank", "acct", "ifsc", "addr",
        "doj", "avsec", "police", "uan", "esic"]


def row(**kw):
    kw.setdefault("name", "Asha Rao")
    kw.setdefault("mail", "asha@example.com")
    return [kw.get(k) for k in KEYS]


def sheet(*rows, header=HEADER, above=0):
    return [["Title row"] + [None] * 14 for _ in range(above)] + [header] + list(rows)


class EngineBase(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user("owner", password="pw-owner-123")
        self.org = Organization.objects.create(name="Acme", created_by=self.owner)
        self.other_org = Organization.objects.create(name="Other", created_by=self.owner)

    def run_rows(self, *rows, **kw):
        return analyze_rows(sheet(*rows, **kw), self.org)


class HeaderAndCleaningTests(EngineBase):
    def test_header_found_below_title_rows_and_case_insensitive(self):
        header = [h.lower() for h in HEADER]
        report = analyze_rows(sheet(row(), header=header, above=2), self.org)
        self.assertTrue(report.ok)
        self.assertEqual(report.rows[0].row_number, 4)  # sheet row numbers

    def test_missing_required_column_is_a_file_error(self):
        report = analyze_rows([["Emp Name", "Phone"], ["A", "1"]], self.org)
        self.assertFalse(report.ok)
        self.assertIn("Mail ID", report.file_errors[0])

    def test_aliases_and_unrecognised_columns(self):
        report = analyze_rows([["Employee Name", "Email", "Pan No", "Weird"], ["Asha Rao", "a@x.com", "abcde1234f", "z"]], self.org)
        self.assertEqual(report.rows[0].data["pan_no"], "ABCDE1234F")
        self.assertIn("Weird", report.notices[0])

    def test_full_row_cleaned(self):
        r = self.run_rows(row(name="  Asha   Devi Rao ", mail="ASHA@Example.com", phone=9876543210.0,
                              aadhar="1234 5678 9012", pan="abcde1234f", ifsc="sbin 0001234",
                              doj="31/03/2026", avsec=datetime(2026, 1, 5), addr="12 MG Road"))
        result = r.rows[0]
        self.assertEqual(result.status, READY)
        self.assertEqual(result.data["first_name"], "Asha")
        self.assertEqual(result.data["last_name"], "Devi Rao")
        self.assertEqual(result.data["email"], "asha@example.com")
        self.assertEqual(result.data["phone"], "9876543210")
        self.assertEqual(result.data["aadhar_no"], "123456789012")
        self.assertEqual(result.data["ifsc_no"], "SBIN0001234")
        self.assertEqual(result.data["date_joined"], date(2026, 3, 31))
        avsec = ChecklistItem.objects.get(organization=self.org, name="AVSEC Training")
        self.assertEqual(result.checklist, {avsec.id: date(2026, 1, 5)})

    def test_required_fields_block_the_row(self):
        for bad in (row(name=None), row(mail=None), row(mail="not-an-email")):
            self.assertEqual(self.run_rows(bad).rows[0].status, ERROR)

    def test_invalid_optional_values_warn_and_are_dropped(self):
        result = self.run_rows(row(aadhar="12345", pan="BAD", doj="not a date", phone="abc")).rows[0]
        self.assertEqual(result.status, WARNING)  # the original importer showed READY here
        self.assertEqual(len(result.warnings), 4)
        for key in ("aadhar_no", "pan_no", "date_joined", "phone"):
            self.assertNotIn(key, result.data)

    def test_aadhar_with_letters_is_rejected_not_silently_stripped(self):
        self.assertEqual(self.run_rows(row(aadhar="ab123456789012")).rows[0].status, WARNING)

    def test_all_supported_date_formats(self):
        for text in ("2026-03-31", "31-03-2026", "31/03/2026", "31.03.2026", "2026.03.31", "31/03/2026 00:00"):
            self.assertEqual(self.run_rows(row(doj=text)).rows[0].data["date_joined"], date(2026, 3, 31), text)

    def test_date_in_a_text_column_is_a_warning(self):
        self.assertEqual(self.run_rows(row(desig=datetime(2026, 1, 1))).rows[0].status, WARNING)

    def test_blank_rows_skipped_and_row_cap(self):
        report = self.run_rows(row(), [None] * 15, row(mail="b@x.com"))
        self.assertEqual(len(report.rows), 2)
        many = [row(mail=f"u{i}@x.com") for i in range(501)]
        self.assertIn("maximum", self.run_rows(*many).file_errors[0])

    def test_custom_and_inactive_checklist_columns(self):
        custom = ChecklistItem.objects.create(organization=self.org, name="Offer letter", order=50)
        ChecklistItem.objects.filter(organization=self.org, name="Police Verification").update(is_active=False)
        header = HEADER + ["Date of Offer letter"]
        report = analyze_rows(sheet(row(police="01/01/2026") + ["02/02/2026"], header=header), self.org)
        self.assertEqual(report.rows[0].checklist, {custom.id: date(2026, 2, 2)})  # police column ignored
        self.assertIn("Date of Police Verification", report.notices[0])


class DuplicateTests(EngineBase):
    def test_duplicate_email_in_file_skips_the_later_row(self):
        report = self.run_rows(row(), row(mail="ASHA@example.com", name="Other"))
        self.assertEqual([r.status for r in report.rows], [READY, SKIP])

    def test_existing_login_and_existing_employee_are_skipped(self):
        User.objects.create_user("login@example.com", email="login@example.com")
        Employee.objects.create(organization=self.org, first_name="Old", email="old@example.com")
        report = self.run_rows(row(mail="login@example.com"), row(mail="old@example.com"), row(mail="new@example.com"))
        self.assertEqual([r.status for r in report.rows], [SKIP, SKIP, READY])

    def test_aadhar_and_pan_clashes_block_but_only_within_the_organization(self):
        Employee.objects.create(organization=self.org, first_name="Old", aadhar_no="123456789012", pan_no="ABCDE1234F")
        Employee.objects.create(organization=self.other_org, first_name="Far", aadhar_no="999999999999")
        report = self.run_rows(
            row(mail="a@x.com", aadhar="123456789012"),
            row(mail="b@x.com", pan="ABCDE1234F"),
            row(mail="c@x.com", aadhar="999999999999"),          # other org -> fine
            row(mail="d@x.com", aadhar="111122223333"),
            row(mail="e@x.com", aadhar="111122223333"),          # clashes with row above in file
        )
        self.assertEqual([r.status for r in report.rows], [ERROR, ERROR, READY, READY, ERROR])

    def test_errored_row_does_not_reserve_its_identifiers(self):
        report = self.run_rows(row(name=None, aadhar="111122223333"), row(mail="b@x.com", aadhar="111122223333"))
        self.assertEqual([r.status for r in report.rows], [ERROR, READY])

    def test_shared_phone_warns(self):
        report = self.run_rows(row(phone="9876543210"), row(mail="b@x.com", phone="+91 98765 43210"))
        self.assertEqual([r.status for r in report.rows], [READY, WARNING])

    def test_query_count_does_not_grow_with_rows(self):
        def queries(n):
            with CaptureQueriesContext(connection) as ctx:
                self.run_rows(*[row(mail=f"u{i}@x.com") for i in range(n)])
            return len(ctx)
        self.assertEqual(queries(3), queries(60))


def xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for excel_row in ws.iter_rows():
        for cell in excel_row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # keep "=..." as literal text, not a formula
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(rows, name="staff.xlsx"):
    return SimpleUploadedFile(name, xlsx_bytes(rows))


class FileReadingTests(EngineBase):
    def test_reads_real_xlsx(self):
        report = parse_upload(upload(sheet(row(doj=datetime(2026, 3, 31)))), self.org)
        self.assertTrue(report.ok)
        self.assertEqual(report.rows[0].data["date_joined"], date(2026, 3, 31))

    def test_garbage_file_is_a_friendly_error_not_an_exception(self):
        report = parse_upload(SimpleUploadedFile("x.xlsx", b"this is not excel"), self.org)
        self.assertIn("valid .xlsx", report.file_errors[0])


class ImportFlowTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user("owner", password="pw-owner-123")
        self.org = Organization.objects.create(name="Acme", created_by=self.owner)
        OrganizationMember.objects.create(organization=self.org, user=self.owner, role="owner")
        self.url = reverse("employee_import", args=[self.org.id])
        self.client.force_login(self.owner)

    def preview(self, *rows):
        return self.client.post(self.url, {"file": upload(sheet(*rows))})

    def confirm(self, response):
        return self.client.post(self.url, {"action": "confirm", "token": response.context["token"]})

    def test_full_flow_creates_employee_login_membership_and_checklist(self):
        response = self.preview(row(phone="9876543210", aadhar="123456789012", avsec="05/01/2026"),
                                row(mail="bad", name="Bad Row"))
        self.assertEqual(response.context["importable"], 1)
        self.assertEqual(response.context["problems"], 1)
        result = self.confirm(response).context["result"]
        self.assertEqual(len(result.created), 1)

        emp = Employee.objects.get(organization=self.org, user__email="asha@example.com")
        self.assertEqual(emp.aadhar_no, "123456789012")
        user = emp.user
        self.assertFalse(user.has_usable_password())
        self.assertEqual(OrganizationMember.objects.get(user=user, organization=self.org).role, "member")
        entry = EmployeeChecklistEntry.objects.get(employee=emp, item__name="AVSEC Training")
        self.assertEqual((entry.status, entry.completed_on, entry.marked_by), ("done", date(2026, 1, 5), self.owner))
        self.assertEqual(Employee.objects.filter(organization=self.org).count(), 2)  # owner + Asha, no duplicates

    def test_imported_employee_is_guided_to_set_a_password(self):
        self.confirm(self.preview(row()))
        self.client.logout()
        response = self.client.post(reverse("login"), {"identifier": "asha@example.com"})
        self.assertRedirects(response, reverse("login_setup_pending"))
        self.client.post(reverse("login_setup_pending"))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("asha@example.com", mail.outbox[0].to)

    def test_confirm_twice_or_with_bad_token_imports_nothing_more(self):
        response = self.preview(row())
        self.confirm(response)
        again = self.confirm(response)
        self.assertEqual(again.status_code, 302)
        bad = self.client.post(self.url, {"action": "confirm", "token": "nope"})
        self.assertEqual(bad.status_code, 302)
        self.assertEqual(Employee.objects.filter(user__email="asha@example.com").count(), 1)

    def test_conflict_created_after_preview_fails_only_that_row(self):
        response = self.preview(row(mail="a@x.com", aadhar="111122223333"), row(mail="b@x.com"))
        Employee.objects.create(organization=self.org, first_name="Sneaky", aadhar_no="111122223333")
        result = self.confirm(response).context["result"]
        self.assertEqual([n for n, *_ in result.created], [3])
        self.assertEqual(len(result.failed), 1)
        self.assertIn("Aadhar", result.failed[0][1])

    def test_session_payload_is_encrypted(self):
        self.preview(row(aadhar="123456789012"))
        blob = self.client.session["employee_import_batch"]["blob"]
        self.assertNotIn("123456789012", blob)
        self.assertNotIn("asha@example.com", blob)

    def test_downloads(self):
        response = self.preview(row(), row(mail="bad"))
        problems = self.client.get(reverse("employee_import_problems", args=[self.org.id]) + f"?token={response.context['token']}")
        ws = load_workbook(BytesIO(problems.content)).active
        self.assertEqual(ws.max_row, 2)  # header + the one bad row
        template = self.client.get(reverse("employee_import_template", args=[self.org.id]))
        headers = [c.value for c in load_workbook(BytesIO(template.content)).worksheets[0][1]]
        self.assertIn("Mail ID", headers)
        self.assertIn("Date of AVSEC Training", headers)

    def test_formula_like_values_are_not_written_as_formulas(self):
        response = self.preview(row(name="=HYPERLINK(\"x\")", mail="bad"))
        problems = self.client.get(reverse("employee_import_problems", args=[self.org.id]) + f"?token={response.context['token']}")
        cell = load_workbook(BytesIO(problems.content)).active["C2"]
        self.assertNotEqual(cell.data_type, "f")  # stored as text, not a formula
        self.assertTrue(cell.value.startswith("=HYPERLINK"))

    def test_only_admins_can_use_it(self):
        for role in ("manager", "member"):
            user = User.objects.create_user(role, password="pw-test-12345")
            OrganizationMember.objects.create(organization=self.org, user=user, role=role)
            self.client.force_login(user)
            for name in ("employee_import", "employee_import_template", "employee_import_problems"):
                self.assertEqual(self.client.get(reverse(name, args=[self.org.id])).status_code, 302, (role, name))
            self.assertEqual(self.client.post(self.url, {"file": upload(sheet(row()))}).status_code, 302)
        self.assertFalse(Employee.objects.filter(user__email="asha@example.com").exists())

    def test_wrong_extension_and_button_visibility(self):
        response = self.client.post(self.url, {"file": SimpleUploadedFile("staff.csv", b"a,b")})
        self.assertContains(response, "Please upload an .xlsx")
        self.assertContains(self.client.get(reverse("employee_list", args=[self.org.id])), "Import employees")


class CrossOrganizationImportTests(TestCase):
    """The same person can be imported into a second organization and keeps ONE login."""

    def setUp(self):
        self.owner_a = User.objects.create_user(
            "owner-a", email="owner-a@example.com", password="pw-owner-123"
        )
        self.org_a = Organization.objects.create(name="Alpha", created_by=self.owner_a)
        OrganizationMember.objects.create(organization=self.org_a, user=self.owner_a, role="owner")

        self.owner_b = User.objects.create_user(
            "owner-b", email="owner-b@example.com", password="pw-owner-123"
        )
        self.org_b = Organization.objects.create(name="Beta", created_by=self.owner_b)
        OrganizationMember.objects.create(organization=self.org_b, user=self.owner_b, role="owner")

    def preview(self, org, owner, *rows):
        self.client.force_login(owner)
        return self.client.post(
            reverse("employee_import", args=[org.id]), {"file": upload(sheet(*rows))}
        )

    def confirm(self, org, response):
        return self.client.post(
            reverse("employee_import", args=[org.id]),
            {"action": "confirm", "token": response.context["token"]},
        ).context["result"]

    def test_same_person_is_linked_into_a_second_organization(self):
        self.confirm(self.org_a, self.preview(self.org_a, self.owner_a, row()))
        asha = User.objects.get(email="asha@example.com")

        response = self.preview(self.org_b, self.owner_b, row(phone="9876543210"))
        self.assertEqual(response.context["counts"]["READY"], 1)
        result = self.confirm(self.org_b, response)

        self.assertEqual(len(result.created), 1)
        self.assertEqual(result.linked, [2])
        self.assertEqual(User.objects.filter(email="asha@example.com").count(), 1)
        employee_b = Employee.objects.get(organization=self.org_b, user=asha)
        self.assertEqual(employee_b.phone, "9876543210")
        self.assertEqual(Employee.objects.filter(user=asha).count(), 2)  # one per organization
        self.assertEqual(
            OrganizationMember.objects.get(organization=self.org_b, user=asha).role, "member"
        )
        self.assertFalse(User.objects.get(pk=asha.pk).has_usable_password())  # untouched

    def test_preview_says_the_login_will_be_linked(self):
        self.confirm(self.org_a, self.preview(self.org_a, self.owner_a, row()))
        response = self.preview(self.org_b, self.owner_b, row())
        shown = response.context["rows"][0]
        self.assertEqual(shown["status"], "READY")
        self.assertIn("linked to this organization", " ".join(shown["notes"]))

    def test_people_already_in_this_organization_are_skipped(self):
        # Org A's owner has an employee record with no email, but is already a member here.
        response = self.preview(
            self.org_a, self.owner_a, row(name="Owner A", mail="owner-a@example.com")
        )
        self.assertEqual(response.context["counts"]["SKIP"], 1)
        self.assertEqual(response.context["importable"], 0)
        self.assertIn("already in this organization", response.context["rows"][0]["notes"][0])

    def test_account_without_any_organization_is_not_linked(self):
        User.objects.create_user(
            "loner@example.com", email="loner@example.com", password="pw-test-12345"
        )
        response = self.preview(self.org_b, self.owner_b, row(mail="loner@example.com"))
        self.assertEqual(response.context["counts"]["SKIP"], 1)
        self.assertIn("any organization", response.context["rows"][0]["notes"][0])

    def test_membership_created_after_preview_fails_only_that_row(self):
        self.confirm(self.org_a, self.preview(self.org_a, self.owner_a, row()))
        asha = User.objects.get(email="asha@example.com")

        response = self.preview(
            self.org_b, self.owner_b, row(), row(mail="other@example.com", name="Other Person")
        )
        # She gets added to org B by someone else between preview and confirm.
        OrganizationMember.objects.create(organization=self.org_b, user=asha, role="member")

        result = self.confirm(self.org_b, response)
        self.assertEqual([n for n, *_ in result.created], [3])
        self.assertEqual(len(result.failed), 1)
        self.assertIn("already in this organization", result.failed[0][1])
        self.assertEqual(Employee.objects.filter(organization=self.org_b, user=asha).count(), 1)

    def test_link_existing_login_refuses_an_account_without_an_organization(self):
        User.objects.create_user(
            "loner@example.com", email="loner@example.com", password="pw-test-12345"
        )
        employee = Employee.objects.create(
            organization=self.org_b, first_name="Loner", email="loner@example.com"
        )
        with self.assertRaises(GrantLoginError):
            link_existing_login(employee)
        employee.refresh_from_db()
        self.assertIsNone(employee.user_id)